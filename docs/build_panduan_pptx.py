"""Generate SeaRIA usage-guide PowerPoint for panitia and peserta."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

TEAL = RGBColor(15, 118, 110)
TEAL_DARK = RGBColor(19, 78, 74)
TEAL_SOFT = RGBColor(204, 251, 241)
SLATE = RGBColor(15, 23, 42)
SLATE_MUTED = RGBColor(71, 85, 105)
WHITE = RGBColor(255, 255, 255)
AMBER = RGBColor(146, 64, 14)
AMBER_BG = RGBColor(255, 251, 235)
CARD = RGBColor(248, 250, 252)
LINE = RGBColor(226, 232, 240)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size=18, bold=False, color=SLATE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_round(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, l, t, w, h, items, size=18, color=SLATE, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)
    return box


def footer(slide, page, total):
    add_rect(slide, 0, Inches(7.22), W, Inches(0.28), TEAL_DARK)
    add_text(slide, Inches(0.5), Inches(7.22), Inches(9), Inches(0.28), "SeaRIA  ·  Panduan pemakaian sistem", 11, False, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(10.5), Inches(7.22), Inches(2.3), Inches(0.28), f"{page} / {total}", 11, False, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def content_header(slide, kicker, title):
    add_rect(slide, 0, 0, W, Inches(1.15), TEAL_DARK)
    add_rect(slide, 0, 0, Inches(0.18), Inches(1.15), TEAL)
    add_text(slide, Inches(0.5), Inches(0.12), Inches(12), Inches(0.32), kicker.upper(), 12, True, TEAL_SOFT)
    add_text(slide, Inches(0.5), Inches(0.42), Inches(12.3), Inches(0.6), title, 28, True, WHITE)


def section_slide(prs, kicker, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, TEAL_DARK)
    add_rect(slide, 0, 0, Inches(0.28), H, TEAL)
    add_text(slide, Inches(0.9), Inches(2.2), Inches(11.5), Inches(0.4), kicker.upper(), 16, True, TEAL_SOFT)
    add_text(slide, Inches(0.9), Inches(2.65), Inches(11.5), Inches(1.3), title, 40, True, WHITE)
    add_text(slide, Inches(0.9), Inches(4.1), Inches(11), Inches(1.4), subtitle, 20, False, TEAL_SOFT)
    return slide


def card(slide, l, t, w, h, title, body, number=None):
    add_round(slide, l, t, w, h, CARD)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(0.12), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if number:
        add_text(slide, l + Inches(0.28), t + Inches(0.12), w - Inches(0.4), Inches(0.28), number, 12, True, TEAL)
        add_text(slide, l + Inches(0.28), t + Inches(0.38), w - Inches(0.45), Inches(0.4), title, 16, True, SLATE)
        add_text(slide, l + Inches(0.28), t + Inches(0.78), w - Inches(0.45), h - Inches(0.9), body, 13, False, SLATE_MUTED)
    else:
        add_text(slide, l + Inches(0.28), t + Inches(0.18), w - Inches(0.45), Inches(0.4), title, 16, True, SLATE)
        add_text(slide, l + Inches(0.28), t + Inches(0.58), w - Inches(0.45), h - Inches(0.7), body, 13, False, SLATE_MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides_meta = []

    def new():
        s = prs.slides.add_slide(blank)
        slides_meta.append(s)
        return s

    # 1 Cover
    s = new()
    add_rect(s, 0, 0, W, H, TEAL_DARK)
    add_rect(s, 0, 0, Inches(0.28), H, TEAL)
    add_text(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.4), "SISTEM INFORMASI KEJUARAAN RENANG", 14, True, TEAL_SOFT)
    add_text(s, Inches(0.9), Inches(2.15), Inches(11.5), Inches(1.6), "Panduan pemakaian SeaRIA", 44, True, WHITE)
    add_text(s, Inches(0.9), Inches(3.85), Inches(11), Inches(0.9), "Untuk panitia dan peserta.\nDari persiapan kejuaraan sampai hasil dipublikasikan.", 20, False, TEAL_SOFT)
    add_text(s, Inches(0.9), Inches(6.3), Inches(11), Inches(0.4), "Tanpa pembayaran  ·  Peserta daftar tanpa akun  ·  Yang login: panitia, super admin, dan juri", 14, False, WHITE)

    # 2 Isi
    s = new()
    content_header(s, "Isi presentasi", "Apa yang akan dibahas")
    items = [
        ("01", "Gambaran sistem", "Peran, alur besar, dan status kejuaraan yang harus berurutan."),
        ("02", "Panduan panitia", "Siapkan acara, terima peserta, bagi seri, hari lomba, publikasi."),
        ("03", "Panduan peserta", "Daftar tanpa akun, isi catatan waktu, simpan kode REG-…"),
        ("04", "Hari lomba & juri", "Siapa mengisi waktu, apa yang panitia cek sebelum hasil terbit."),
    ]
    for i, (num, title, body) in enumerate(items):
        x = Inches(0.5) + (i % 2) * Inches(6.35)
        y = Inches(1.5) + (i // 2) * Inches(2.55)
        card(s, x, y, Inches(6.1), Inches(2.3), title, body, num)

    # 3 Apa itu
    s = new()
    content_header(s, "Pengenalan", "Apa itu SeaRIA?")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.3), Inches(0.7), "Satu tempat untuk pendaftaran, pembagian seri, buku acara, dan hasil resmi kejuaraan renang.", 20, False, SLATE_MUTED)
    points = [
        ("Tanpa bayar di sistem", "Pendaftaran cukup data atlet dan nomor lomba. Tidak ada tagihan di aplikasi."),
        ("Peserta tanpa akun", "Orang tua / pelatih / atlet mengisi form publik. Tidak perlu daftar akun."),
        ("Panitia yang mengendalikan", "Panitia menyiapkan acara, memeriksa data, membagi seri, lalu menerbitkan hasil."),
        ("Urutan wajib", "Status kejuaraan maju selangkah demi selangkah. Tidak boleh loncat."),
    ]
    for i, (title, body) in enumerate(points):
        x = Inches(0.5) + (i % 2) * Inches(6.35)
        y = Inches(2.25) + (i // 2) * Inches(2.2)
        card(s, x, y, Inches(6.1), Inches(2.0), title, body)

    # 4 Peran
    s = new()
    content_header(s, "Peran", "Siapa mengerjakan apa")
    roles = [
        ("Panitia", "Siapkan kejuaraan, buka/tutup pendaftaran, verifikasi, bagi seri, buku acara, cek hasil, publikasi."),
        ("Peserta", "Daftar tanpa akun di Beranda atau /daftar. Simpan kode REG-…. Koreksi hanya lewat panitia."),
        ("Juri", "Isi waktu di hari lomba sesuai tugas. Kunci seri setelah lengkap."),
        ("Super Admin", "Sama seperti panitia, plus boleh mundurkan status kejuaraan (wajib isi alasan)."),
    ]
    for i, (title, body) in enumerate(roles):
        y = Inches(1.45) + i * Inches(1.3)
        add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.15), CARD)
        add_rect(s, Inches(0.5), y, Inches(0.14), Inches(1.15), TEAL)
        add_text(s, Inches(0.9), y + Inches(0.18), Inches(2.4), Inches(0.8), title, 20, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(3.5), y + Inches(0.22), Inches(8.9), Inches(0.75), body, 16, False, SLATE_MUTED, anchor=MSO_ANCHOR.MIDDLE)

    # 5 Alur
    s = new()
    content_header(s, "Alur besar", "10 langkah kejuaraan")
    steps = [
        "Siapkan data, kelompok umur, nomor, matriks",
        "Buka pendaftaran",
        "Peserta daftar / panitia unggah Excel",
        "Panitia setujui entri",
        "Tutup pendaftaran",
        "Bagi seri dan lintasan",
        "Cetak buku acara",
        "Hari lomba: isi hasil",
        "Panitia cek hasil",
        "Publikasi ke publik",
    ]
    for i, text in enumerate(steps):
        col = i % 5
        row = i // 5
        x = Inches(0.4) + col * Inches(2.55)
        y = Inches(1.55) + row * Inches(2.55)
        add_round(s, x, y, Inches(2.4), Inches(2.25), CARD)
        add_text(s, x + Inches(0.15), y + Inches(0.25), Inches(2.1), Inches(0.55), f"{i+1:02d}", 28, True, TEAL, PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.12), y + Inches(0.9), Inches(2.16), Inches(1.15), text, 14, False, SLATE, PP_ALIGN.CENTER)

    # 6 Status
    s = new()
    content_header(s, "Status kejuaraan", "Harus maju berurutan — tidak boleh loncat")
    statuses = ["Draf", "Pendaftaran terbuka", "Pendaftaran ditutup", "Sudah diseeding", "Hari lomba", "Selesai", "Dipublikasikan"]
    for i, name in enumerate(statuses):
        y = Inches(1.45) + i * Inches(0.68)
        add_round(s, Inches(0.5), y, Inches(8.6), Inches(0.58), CARD)
        add_text(s, Inches(0.7), y, Inches(0.6), Inches(0.58), str(i + 1), 16, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(1.4), y, Inches(7.5), Inches(0.58), name, 18, True, SLATE, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(statuses) - 1:
            add_text(s, Inches(9.2), y, Inches(0.5), Inches(0.58), "↓", 18, True, TEAL, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_round(s, Inches(9.7), Inches(1.45), Inches(3.1), Inches(4.7), AMBER_BG)
    add_text(s, Inches(9.95), Inches(1.7), Inches(2.65), Inches(0.4), "Ingat", 16, True, AMBER)
    add_text(s, Inches(9.95), Inches(2.2), Inches(2.65), Inches(3.6), "Mundur status hanya Super Admin, dan wajib mengisi alasan.\n\nForm publik tidak muncul selama masih Draf.", 15, False, AMBER)

    # Section panitia
    section_slide(prs, "Bagian 1", "Panduan panitia", "Siapkan dulu, baru peserta bisa daftar. Panitia mengendalikan seluruh alur sampai hasil terbit.")
    slides_meta.append(prs.slides[-1])

    # 8 Menu
    s = new()
    content_header(s, "Panitia", "Menu di sidebar kiri")
    menus = [
        ("Dasbor", "Pilih atau buat acara. Ubah status. Pengaturan kelompok umur, nomor, matriks, kesiapan, juri."),
        ("Pendaftaran", "Import Excel, tambah manual, setujui atau tolak entri."),
        ("Pembagian seri", "Susun seri dan lintasan dari catatan waktu. Cek, kunci."),
        ("Buku acara", "Lihat start list, unduh PDF, unduh lembar hasil kosong."),
        ("Hasil", "Isi/koreksi waktu, verifikasi, medali, export."),
        ("Klub / Atlet", "Data master. Verifikasi klub baru yang diketik peserta."),
    ]
    for i, (title, body) in enumerate(menus):
        x = Inches(0.45) + (i % 3) * Inches(4.2)
        y = Inches(1.5) + (i // 3) * Inches(2.55)
        card(s, x, y, Inches(4.0), Inches(2.3), title, body)

    # 9 Siapkan
    s = new()
    content_header(s, "Panitia · sebelum daftar dibuka", "Siapkan kejuaraan")
    prep = [
        ("1. Klub & atlet", "Opsional. Tambah klub yang sudah dikenal. Verifikasi klub yang masih menunggu."),
        ("2. Buat acara", "Dasbor → Tambah acara. Nama, tempat, tanggal, 6/8 lintasan, batas nomor per atlet."),
        ("3. Kelompok umur", "Berdasarkan tahun lahir, bukan umur kalender. Rentang tidak boleh tumpang tindih."),
        ("4. Nomor lomba", "Bisa isi susunan baku 34 nomor PA/PI. Contoh: 13 = 50 m gaya dada putra."),
        ("5. Matriks", "Centang grup mana boleh ikut nomor mana. Form daftar hanya menampilkan yang diizinkan."),
        ("6. Kesiapan", "Cek yang masih kurang, lalu lanjut status ke Pendaftaran terbuka."),
    ]
    for i, (title, body) in enumerate(prep):
        x = Inches(0.45) + (i % 3) * Inches(4.2)
        y = Inches(1.45) + (i // 3) * Inches(2.6)
        card(s, x, y, Inches(4.0), Inches(2.4), title, body)

    # 10 Terima peserta
    s = new()
    content_header(s, "Panitia · pendaftaran terbuka", "Tiga cara data peserta masuk")
    ways = [
        ("Form publik", "Peserta buka Beranda atau /daftar. Tanpa akun. Dapat kode REG-…"),
        ("Import Excel", "Panitia & Super Admin. Pendaftaran → Import Excel. Unduh template, isi PESERTA, unggah. NOMOR LOMBA hanya rujukan terkunci."),
        ("Input manual", "Pendaftaran → Tambah manual. Untuk koreksi kecil, bukan daftar klub utuh."),
    ]
    for i, (title, body) in enumerate(ways):
        card(s, Inches(0.5) + i * Inches(4.2), Inches(1.55), Inches(4.0), Inches(3.3), title, body, f"0{i+1}")
    add_round(s, Inches(0.5), Inches(5.1), Inches(12.3), Inches(1.7), AMBER_BG)
    add_text(s, Inches(0.75), Inches(5.25), Inches(11.8), Inches(0.35), "Penting", 16, True, AMBER)
    add_text(s, Inches(0.75), Inches(5.6), Inches(11.8), Inches(1.0), "Ketiga jalur menghasilkan entri berstatus menunggu verifikasi. Yang belum disetujui tidak masuk pembagian seri. Import hanya bisa selama status masih Pendaftaran terbuka.", 16, False, AMBER)

    # 11 Import
    s = new()
    content_header(s, "Panitia & Super Admin · Pendaftaran → Import Excel", "Cara unggah daftar klub")
    add_bullets(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(5.3), [
        "Hanya panitia dan Super Admin yang mengunggah. Peserta memakai form publik.",
        "Unduh template. Isi lembar PESERTA. Lembar NOMOR LOMBA terkunci (kode, nama, gender, grup).",
        "Salin KODE ACARA dari NOMOR LOMBA. Jangan ubah GRUP YANG BOLEH IKUT di Excel.",
        "Satu baris = satu atlet pada satu nomor. Ikut tiga nomor = tiga baris. Catatan waktu boleh kosong (NT).",
        "Unggah .xlsx atau .csv (maks. 5 MB, 2.000 baris). Cek pratinjau, perbaiki baris bermasalah.",
        "Tekan Import … baris valid. Lalu tetap verifikasi sebelum pembagian seri.",
    ], 18, SLATE, 10)

    # 12 Verifikasi
    s = new()
    content_header(s, "Panitia · Pendaftaran → Antrean verifikasi", "Setujui atau tolak entri")
    card(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(3.4), "Setujui", "Status menjadi terverifikasi. Hanya entri ini yang masuk pembagian seri. Ada penyetujuan massal. Bisa saring per klub, nomor, atau kelompok umur.", "01")
    card(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(3.4), "Tolak", "Wajib isi alasan. Pendaftar tidak mengedit sendiri. Panitia yang memperbaiki, atau minta data ulang lewat WhatsApp dengan kode REG-…", "02")
    add_round(s, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.65), TEAL_SOFT)
    add_text(s, Inches(0.75), Inches(5.35), Inches(11.8), Inches(1.25), "Jika semua data sudah rapi: buka Ringkasan acara, lanjutkan status ke Pendaftaran ditutup. Form publik dan import Excel berhenti.", 18, False, TEAL_DARK)

    # 13 Seeding artinya
    s = new()
    content_header(s, "Panitia · Pembagian seri", "Apa itu seeding?")
    add_text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(1.0), "Seeding = menyusun siapa berenang di seri berapa dan lintasan berapa, dari catatan waktu saat daftar — bukan hasil lomba.", 20, False, SLATE)
    facts = [
        ("Seri", "Kolam hanya 6 atau 8 lintasan. Jika peserta lebih banyak, mereka dibagi ke beberapa gelombang (seri)."),
        ("Lintasan tengah", "Biasanya untuk yang lebih cepat dalam seri itu. Seri terakhir biasanya berisi perenang tercepat."),
        ("NT", "Tidak ada catatan waktu. Diletakkan di seri belakang."),
        ("Hanya yang disetujui", "Entri menunggu verifikasi tidak ikut. Pembagian dihitung per nomor lomba × kelompok umur."),
    ]
    for i, (title, body) in enumerate(facts):
        x = Inches(0.5) + (i % 2) * Inches(6.35)
        y = Inches(2.5) + (i // 2) * Inches(2.05)
        card(s, x, y, Inches(6.1), Inches(1.9), title, body)

    # 14 Langkah seeding
    s = new()
    content_header(s, "Panitia · Pembagian seri", "Langkah yang harus dikerjakan")
    seed_steps = [
        "Tutup pendaftaran di Ringkasan.",
        "Setujui semua entri di Pendaftaran.",
        "Tekan Bagi seri seluruh kejuaraan.",
        "Buka Lihat susunan per nomor × kelompok umur. Tukar lintasan atau keluarkan peserta jika perlu.",
        "Kunci jika susunan sudah final.",
        "Di Ringkasan, lanjutkan status ke Sudah diseeding. Baru buku acara bisa dicetak dan dilihat publik.",
    ]
    for i, text in enumerate(seed_steps):
        y = Inches(1.4) + i * Inches(0.88)
        add_round(s, Inches(0.5), y, Inches(12.3), Inches(0.78), CARD)
        add_text(s, Inches(0.7), y, Inches(0.7), Inches(0.78), f"{i+1}", 20, True, TEAL, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(1.5), y, Inches(11), Inches(0.78), text, 18, False, SLATE, anchor=MSO_ANCHOR.MIDDLE)

    # 15 Buku acara & juri
    s = new()
    content_header(s, "Panitia · setelah seri dibagi", "Buku acara dan penugasan juri")
    card(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(4.9), "Buku acara", "Menu Buku acara.\n\n• Lihat start list di layar\n• Unduh PDF buku acara\n• Unduh lembar hasil kosong untuk dicatat di pinggir kolam\n\nPublik juga bisa melihat start list setelah status Sudah diseeding.", "01")
    card(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(4.9), "Penugasan juri", "Ringkasan → Pengaturan acara → Penugasan juri.\n\nTentukan juri mana mengisi nomor/seri mana.\n\nTanpa ini, juri tidak melihat tugas. Bisa juga diatur dari halaman Hasil.", "02")

    # 16 Hari lomba
    s = new()
    content_header(s, "Panitia · hari lomba", "Hasil diisi, lalu dicek")
    add_bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(3.4), [
        "Di Ringkasan, lanjutkan status ke Hari lomba.",
        "Juri masuk → Tugas juri → pilih seri → isi waktu per lintasan.",
        "Ketikan cepat: 3470 menjadi 00:34.70. Status khusus: DNS, DNF, DSQ.",
        "Setelah seri lengkap, juri kunci seri.",
        "Panitia boleh mengisi atau mengkoreksi di menu Hasil tanpa menunggu juri. Koreksi setelah kunci tercatat di Audit.",
        "Tab Verifikasi: tandai hasil per seri atau per nomor sudah dicek.",
        "Jika seluruh hasil selesai, lanjutkan status ke Selesai, lalu Dipublikasikan.",
    ], 17, SLATE, 8)

    # 17 Publikasi
    s = new()
    content_header(s, "Panitia · publikasi", "Setelah status Dipublikasikan")
    pubs = [
        ("Hasil resmi", "Publik lihat dan unduh PDF hasil di halaman kejuaraan."),
        ("Peringkat", "Per nomor lomba × kelompok umur."),
        ("Medali / klasemen", "Jika diaktifkan di halaman hasil."),
        ("Cari atlet & arsip", "Menu Cari atlet dan Arsip di situs publik."),
        ("Export", "Unduh Excel peserta, start list, hasil, medali."),
        ("Ingat", "Catatan waktu daftar bukan hasil. Hasil hanya dari hari lomba."),
    ]
    for i, (title, body) in enumerate(pubs):
        x = Inches(0.45) + (i % 3) * Inches(4.2)
        y = Inches(1.5) + (i // 3) * Inches(2.55)
        card(s, x, y, Inches(4.0), Inches(2.35), title, body)

    # Section peserta
    section_slide(prs, "Bagian 2", "Panduan peserta", "Tidak perlu membuat akun. Isi data, pilih nomor, simpan kode pendaftaran.")
    slides_meta.append(prs.slides[-1])

    # 19 3 langkah
    s = new()
    content_header(s, "Peserta", "Daftar dalam 3 langkah")
    card(s, Inches(0.45), Inches(1.5), Inches(4.0), Inches(4.9), "1. Data diri", "Kontak pendaftar (nama + WhatsApp). Nama atlet, L/P, tahun lahir, klub, kota.\n\nTahun lahir menentukan kelompok umur — bukan umur di hari lomba.", "Langkah 1")
    card(s, Inches(4.65), Inches(1.5), Inches(4.0), Inches(4.9), "2. Pilih nomor", "Hanya nomor yang cocok dengan jenis kelamin dan kelompok umur yang tampil.\n\nCentang nomor, isi catatan waktu jika ada. Kosong = NT.", "Langkah 2")
    card(s, Inches(8.85), Inches(1.5), Inches(4.0), Inches(4.9), "3. Cek & kirim", "Periksa data. Setelah dikirim tidak bisa diubah sendiri.\n\nSimpan atau potret kode REG-…. Itu rujukan saat menghubungi panitia.", "Langkah 3")

    # 20 Seed time
    s = new()
    content_header(s, "Peserta · catatan waktu", "Ketik 6 angka, dari kiri ke kanan")
    add_text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.85), "Catatan waktu (seed) = waktu terbaik atlet. Dipakai panitia untuk membagi seri dan lintasan. Bukan hasil lomba nanti.", 18, False, SLATE)
    boxes = [("0", "1", "Menit", "paling kiri"), ("3", "4", "Detik", "tengah"), ("7", "0", "1/100 dtk", "paling kanan")]
    x0 = Inches(2.1)
    for i, (a, b, label, hint) in enumerate(boxes):
        x = x0 + i * Inches(3.1)
        add_round(s, x, Inches(2.45), Inches(1.15), Inches(1.35), TEAL_SOFT)
        add_round(s, x + Inches(1.3), Inches(2.45), Inches(1.15), Inches(1.35), TEAL_SOFT)
        add_text(s, x, Inches(2.55), Inches(1.15), Inches(1.15), a, 36, True, TEAL_DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(1.3), Inches(2.55), Inches(1.15), Inches(1.15), b, 36, True, TEAL_DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(s, x, Inches(3.9), Inches(2.45), Inches(0.35), label, 16, True, TEAL, PP_ALIGN.CENTER)
        add_text(s, x, Inches(4.25), Inches(2.45), Inches(0.3), hint, 13, False, SLATE_MUTED, PP_ALIGN.CENTER)
        if i < 2:
            add_text(s, x + Inches(2.45), Inches(2.7), Inches(0.55), Inches(1.0), ":" if i == 0 else ".", 32, True, SLATE_MUTED, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(0.4), "013470  →  01:34.70   (1 menit 34,70 detik)", 20, True, TEAL_DARK, PP_ALIGN.CENTER)
    add_bullets(s, Inches(1.2), Inches(5.2), Inches(11), Inches(1.6), [
        "Contoh 52,20 detik: ketik 005220 atau 5220.",
        "Boleh juga 52.20 atau 1:34.70. Kosong atau NT = belum punya catatan waktu.",
    ], 16, SLATE_MUTED, 6)

    # 21 Setelah kirim
    s = new()
    content_header(s, "Peserta", "Setelah pendaftaran terkirim")
    after = [
        ("Simpan kode REG-…", "Potret halaman. Kode itu satu-satunya rujukan saat menghubungi panitia. Halaman tidak bisa dibuka lagi setelah peramban ditutup."),
        ("Panitia memeriksa", "Jika ada yang perlu diperbaiki, panitia menghubungi WhatsApp Anda."),
        ("Tidak bisa edit sendiri", "Perubahan atau pembatalan hanya lewat panitia, dengan menyebutkan kode."),
        ("Lihat jadwal & hasil", "Jadwal di halaman kejuaraan. Buku acara setelah diseeding. Hasil setelah dipublikasikan. Cari atlet di menu publik."),
    ]
    for i, (title, body) in enumerate(after):
        x = Inches(0.5) + (i % 2) * Inches(6.35)
        y = Inches(1.5) + (i // 2) * Inches(2.55)
        card(s, x, y, Inches(6.1), Inches(2.35), title, body)

    # Section juri
    section_slide(prs, "Bagian 3", "Hari lomba — juri", "Juri mengisi waktu. Panitia boleh membantu atau mengoreksi.")
    slides_meta.append(prs.slides[-1])

    s = new()
    content_header(s, "Juri", "Tugas di hari lomba")
    add_bullets(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(5.2), [
        "Masuk lewat /login. Langsung ke Tugas juri — hanya seri yang ditugaskan panitia yang tampil.",
        "Pilih seri, isi waktu per lintasan. Ketikan cepat 4–6 angka: 3470 → 00:34.70.",
        "Jika tidak start / tidak finish / didiskualifikasi: pilih DNS, DNF, atau DSQ — jangan isi waktu.",
        "Setelah semua lintasan terisi, kunci seri.",
        "Jika tidak ada tugas: minta panitia mengisi Penugasan juri di pengaturan acara.",
        "Panitia bisa mengisi hasil yang sama dari menu Hasil, termasuk koreksi setelah kunci (tercatat di Audit).",
    ], 18, SLATE, 10)

    # Siapa mengerjakan
    s = new()
    content_header(s, "Ringkasan", "Siapa mengerjakan apa")
    rows = [
        ("Siapkan kejuaraan, grup, nomor, matriks", "Panitia"),
        ("Buka / tutup pendaftaran, bagi seri, publikasi", "Panitia"),
        ("Isi form /daftar", "Peserta (tanpa akun)"),
        ("Import Excel / input manual", "Panitia & Super Admin"),
        ("Setujui atau tolak entri", "Panitia"),
        ("Isi waktu di hari lomba", "Juri, atau panitia"),
        ("Mundurkan status kejuaraan", "Super Admin saja"),
    ]
    add_round(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.5), TEAL)
    add_text(s, Inches(0.7), Inches(1.4), Inches(8.2), Inches(0.5), "Langkah", 14, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(9.0), Inches(1.4), Inches(3.5), Inches(0.5), "Siapa", 14, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    for i, (left, right) in enumerate(rows):
        y = Inches(1.95) + i * Inches(0.65)
        add_round(s, Inches(0.5), y, Inches(12.3), Inches(0.58), CARD if i % 2 == 0 else WHITE)
        add_text(s, Inches(0.7), y, Inches(8.2), Inches(0.58), left, 16, False, SLATE, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(9.0), y, Inches(3.5), Inches(0.58), right, 16, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)

    # Sering terlewat
    s = new()
    content_header(s, "Pengingat", "Yang sering terlewat")
    misses = [
        "Form publik tidak muncul selama kejuaraan masih Draf. Buka status Pendaftaran terbuka dulu.",
        "Peserta yang belum disetujui tidak masuk pembagian seri.",
        "Pembagian seri dihitung per nomor lomba × kelompok umur, bukan per nomor saja.",
        "Catatan waktu daftar bukan hasil lomba. Hasil diisi juri (atau panitia) di hari lomba.",
        "Satu atlet satu kali per nomor. Ikut tiga nomor = tiga baris (form, Excel, maupun manual).",
        "Setelah dikirim, pendaftar tidak mengedit lagi. Koreksi di sisi panitia.",
        "Import Excel hanya selama Pendaftaran terbuka. Setelah ditutup, import berhenti.",
        "Tanpa penugasan juri, juri tidak melihat seri yang harus diisi.",
    ]
    add_bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(5.4), misses, 17, SLATE, 8)

    # Penutup
    s = new()
    add_rect(s, 0, 0, W, H, TEAL_DARK)
    add_rect(s, 0, 0, Inches(0.28), H, TEAL)
    add_text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.2), "Siap menjalankan kejuaraan", 36, True, WHITE)
    add_text(s, Inches(0.9), Inches(3.6), Inches(11.2), Inches(1.4), "Panitia menyiapkan dan mengendalikan alur.\nPeserta daftar tanpa akun, lalu simpan kode REG-…\nJuri mengisi waktu di seri yang ditugaskan.", 20, False, TEAL_SOFT)
    add_text(s, Inches(0.9), Inches(5.5), Inches(11), Inches(0.6), "Masuk panitia / juri: /login     ·     Daftar peserta: /daftar", 16, False, WHITE)

    # Footers on content slides (skip cover, section, closing which are full-bleed)
    total = len(prs.slides)
    for idx, slide in enumerate(prs.slides, start=1):
        fill = slide.shapes[0].fill if slide.shapes else None
        # Always add footer except full teal section/cover/closing: detect by first shape color
        try:
            first = slide.shapes[0]
            rgb = first.fill.fore_color.rgb
            if rgb == TEAL_DARK and first.height >= Inches(7.0):
                continue
        except Exception:
            pass
        footer(slide, idx, total)

    out = r"R:\laragon\www\searia\app\docs\panduan-pemakaian-searia.pptx"
    prs.save(out)
    print(out)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
